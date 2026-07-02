import importlib.util
import json
import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
BUILD_SCRIPT = ROOT / "scripts" / "build_minutes.py"
VALIDATE_SCRIPT = ROOT / "scripts" / "validate_minutes.py"
TEMPLATE = ROOT / "assets" / "materials-department-minutes-template.docx"
SKILL = ROOT / "SKILL.md"
OPENAI_YAML = ROOT / "agents" / "openai.yaml"


def load_build_module():
    if not BUILD_SCRIPT.exists():
        return None
    spec = importlib.util.spec_from_file_location("build_minutes", BUILD_SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class MeetingMinutesSkillTests(unittest.TestCase):
    def test_required_skill_files_exist(self):
        for path in (BUILD_SCRIPT, VALIDATE_SCRIPT, TEMPLATE, SKILL, OPENAI_YAML):
            self.assertTrue(path.exists(), f"missing required skill file: {path}")

    def test_builds_docx_with_fixed_recorder_date_and_plain_table(self):
        module = load_build_module()
        if module is None:
            self.fail("build_minutes.py is not implemented")

        content = {
            "location": "实验楼B515会议室",
            "participants": ["郭源杰", "李豹"],
            "topics": ["同位素材料", "叔丁胺合成"],
            "sections": [
                {"title": "同位素材料", "paragraphs": ["讨论了样品测试进展。"]},
                {"title": "叔丁胺合成", "paragraphs": ["继续优化第一步温度控制。"]},
            ],
            "conclusions": ["先完成重复实验，再评估放大。"],
            "actions": [
                {
                    "item": "完成重复实验",
                    "owner": "郭源杰",
                    "deadline": "待确认",
                    "notes": "复核纯度",
                }
            ],
            "review_items": ["参会人员名单", "关键数值", "截止日期"],
            "source_note": "测试逐字稿",
        }

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "minutes.docx"
            module.build_minutes(content, output, meeting_date=date(2026, 7, 2))
            doc = Document(output)

        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("会议时间：2026年7月2日", text)
        self.assertIn("纪要人员：郭源杰", text)
        self.assertIn("同位素材料", text)
        self.assertIn("叔丁胺合成", text)

        table = doc.tables[0]
        self.assertEqual(
            [cell.text for cell in table.rows[0].cells],
            ["序号", "跟进事宜", "负责人", "截止日期", "备注"],
        )
        self.assertEqual(
            [int(node.get(qn("w:w"))) for node in table._tbl.tblGrid],
            [817, 4527, 1023, 1295, 1580],
        )
        for row in table.rows:
            for cell in row.cells:
                shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
                self.assertTrue(
                    shading is None
                    or shading.get(qn("w:fill")) in (None, "auto", "FFFFFF"),
                    "table cells must not have a fill color",
                )

    def test_rejects_missing_topics(self):
        module = load_build_module()
        if module is None:
            self.fail("build_minutes.py is not implemented")
        content = {
            "topics": [],
            "sections": [],
            "conclusions": [],
            "actions": [],
            "review_items": [],
        }
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "主题"):
                module.build_minutes(content, Path(tmp) / "minutes.docx")

    def test_template_is_clean_and_contains_no_historical_media(self):
        if not TEMPLATE.exists():
            self.fail("clean template asset is not implemented")
        with zipfile.ZipFile(TEMPLATE) as archive:
            names = archive.namelist()
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertFalse(any(name.startswith("word/media/") for name in names))
        self.assertNotIn("2026年5月11日", document_xml)
        self.assertNotIn("叔丁胺合成方案优化", document_xml)


if __name__ == "__main__":
    unittest.main()

