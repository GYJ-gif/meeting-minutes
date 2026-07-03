import importlib.util
import tempfile
import unittest
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "build_minutes.py"
TEMPLATE = ROOT / "assets" / "materials-department-minutes-template.docx"
HEADINGS = [
    "一、会议基本信息",
    "二、会议主要内容",
    "三、会议核心结论（如有）",
    "四、跟进事宜及节点",
    "五、交流图片",
]


def load_module():
    spec = importlib.util.spec_from_file_location("build_minutes_fidelity", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def content():
    return {
        "recorder": "郭源杰",
        "participants": ["郭源杰", "李四"],
        "location": "会议室",
        "topics": ["同位素材料"],
        "sections": [{"title": "同位素材料", "paragraphs": ["讨论测试进展。"]}],
        "conclusions": ["继续推进实验。"],
        "actions": [
            {"item": f"跟进事项{i}", "owner": "郭源杰", "deadline": "待确认", "notes": ""}
            for i in range(1, 7)
        ],
        "review_items": ["专业名称", "关键数值", "截止日期"],
    }


def child_value(parent, tag, attribute):
    child = parent.find(qn(tag)) if parent is not None else None
    return child.get(qn(attribute)) if child is not None else None


class TemplateFidelityTests(unittest.TestCase):
    def build_doc(self):
        temporary = tempfile.TemporaryDirectory()
        output = Path(temporary.name) / "minutes.docx"
        load_module().build_minutes(content(), output, meeting_date=date(2026, 7, 2))
        return temporary, Document(output)

    def test_template_asset_retains_formatted_skeleton(self):
        doc = Document(TEMPLATE)
        texts = [paragraph.text for paragraph in doc.paragraphs]
        for heading in HEADINGS:
            self.assertIn(heading, texts)
        self.assertEqual(len(doc.tables), 1)

    def test_output_has_all_five_blue_headings(self):
        temporary, doc = self.build_doc()
        self.addCleanup(temporary.cleanup)
        headings = {paragraph.text: paragraph for paragraph in doc.paragraphs if paragraph.text in HEADINGS}
        self.assertEqual(list(headings), HEADINGS)
        for text, paragraph in headings.items():
            self.assertTrue(paragraph.runs, f"heading has no runs: {text}")
            for run in paragraph.runs:
                self.assertEqual(child_value(run._element.rPr, "w:color", "w:val"), "2C569A")
                self.assertTrue(run.bold)
                self.assertEqual(run.font.size.pt, 12)

    def test_fifth_heading_is_explicit_and_not_automatically_numbered(self):
        temporary, doc = self.build_doc()
        self.addCleanup(temporary.cleanup)
        fifth = next(paragraph for paragraph in doc.paragraphs if paragraph.text == "五、交流图片")
        num_pr = fifth._p.pPr.find(qn("w:numPr")) if fifth._p.pPr is not None else None
        self.assertIsNone(num_pr)

    def test_output_table_matches_template_format(self):
        temporary, doc = self.build_doc()
        self.addCleanup(temporary.cleanup)
        table = doc.tables[0]
        properties = table._tbl.tblPr
        borders = properties.find(qn("w:tblBorders"))
        margins = properties.find(qn("w:tblCellMar"))
        self.assertEqual(child_value(properties, "w:tblLayout", "w:type"), "fixed")
        self.assertEqual(child_value(properties, "w:tblInd", "w:w"), "0")
        self.assertEqual(
            {name: child_value(borders, f"w:{name}", "w:val") for name in ("top", "left", "bottom", "right", "insideH", "insideV")},
            {"top": "single", "left": "none", "bottom": "single", "right": "none", "insideH": "single", "insideV": "single"},
        )
        self.assertEqual(
            {name: child_value(margins, f"w:{name}", "w:w") for name in ("top", "left", "bottom", "right")},
            {"top": "0", "left": "108", "bottom": "0", "right": "108"},
        )
        self.assertEqual([node.get(qn("w:w")) for node in table._tbl.tblGrid], ["817", "4527", "1023", "1295", "1580"])
        self.assertEqual(len(table.rows), 7)
        for row in table.rows:
            for cell in row.cells:
                self.assertIsNone(cell.vertical_alignment)
                self.assertEqual(cell.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
                shading = cell._tc.tcPr.find(qn("w:shd"))
                fill = shading.get(qn("w:fill")) if shading is not None else None
                self.assertIn(fill, (None, "auto", "FFFFFF"))


if __name__ == "__main__":
    unittest.main()
