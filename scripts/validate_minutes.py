#!/usr/bin/env python3
"""Validate a generated meeting-minutes DOCX."""
from __future__ import annotations

import argparse
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from build_minutes import TABLE_GRID_DXA, TABLE_HEADERS, build_title, today_shanghai


HEADING_TEXTS = [
    "一、会议基本信息",
    "二、会议主要内容",
    "三、会议核心结论（如有）",
    "四、跟进事宜及节点",
    "五、交流图片",
]
HEADING_COLOR = "2C569A"


def _child_value(parent, tag, attribute):
    child = parent.find(qn(tag)) if parent is not None else None
    return child.get(qn(attribute)) if child is not None else None


def validate_minutes(path, topics=None, recorder=None, participants=None):
    path = Path(path).resolve()
    issues = []
    if not path.is_file():
        return [f"文件不存在：{path}"]
    try:
        with zipfile.ZipFile(path) as archive:
            bad = archive.testzip()
            if bad:
                issues.append(f"DOCX 压缩包损坏：{bad}")
    except zipfile.BadZipFile:
        return ["文件不是有效的 DOCX/ZIP 包"]
    try:
        doc = Document(path)
    except Exception as exc:
        return [f"python-docx 无法打开文件：{exc}"]

    paragraphs = [paragraph.text for paragraph in doc.paragraphs]
    text = "\n".join(paragraphs)
    current = today_shanghai()
    expected_date = f"会议时间：{current.year}年{current.month}月{current.day}日"
    if expected_date not in text:
        issues.append(f"会议日期不是北京时间当天：应包含“{expected_date}”")

    topics = [str(item).strip() for item in topics or [] if str(item).strip()]
    if topics and build_title(topics) not in text:
        issues.append(f"标题未按主题生成：{build_title(topics)}")
    for topic in topics:
        if topic not in text:
            issues.append(f"未覆盖会议主题：{topic}")

    recorder_line = next((item for item in paragraphs if item.startswith("纪要人员：")), "")
    if recorder and recorder_line != f"纪要人员：{recorder.strip()}":
        issues.append(f"纪要人员不匹配：{recorder}")
    if not recorder and recorder_line in ("", "纪要人员：", "纪要人员：待核验", "纪要人员：待确认"):
        issues.append("纪要人员为空或未确认")

    participant_line = next((item for item in paragraphs if item.startswith("参与人员：")), "")
    expected_people = "、".join(str(item).strip() for item in participants or [] if str(item).strip())
    if expected_people and participant_line != f"参与人员：{expected_people}":
        issues.append(f"参会人员不匹配：{expected_people}")
    if not expected_people and participant_line in ("", "参与人员：", "参与人员：待核验", "参与人员：待确认"):
        issues.append("参会人员为空或未确认")

    for item in HEADING_TEXTS + [
        "下次交流预计时间：",
        "预计议题：",
        "如涉及比较好的学习材料和要点可分享给大家共同学习",
    ]:
        if item not in text:
            issues.append(f"缺少模板固定内容：{item}")

    for heading_text in HEADING_TEXTS:
        paragraph = next((item for item in doc.paragraphs if item.text == heading_text), None)
        if paragraph is None:
            continue
        for run in paragraph.runs:
            if _child_value(run._element.rPr, "w:color", "w:val") != HEADING_COLOR:
                issues.append(f"章节标题颜色不匹配：{heading_text}")
                break
            if run.bold is not True or run.font.size is None or run.font.size.pt != 12:
                issues.append(f"章节标题字号或加粗不匹配：{heading_text}")
                break

    fifth = next((item for item in doc.paragraphs if item.text == "五、交流图片"), None)
    if fifth is not None and fifth._p.pPr is not None and fifth._p.pPr.find(qn("w:numPr")) is not None:
        issues.append("第五标题不得保留自动编号，避免重复显示")

    if len(doc.tables) != 1:
        issues.append(f"跟进事项表数量应为1，实际为{len(doc.tables)}")
        return issues

    table = doc.tables[0]
    if [cell.text for cell in table.rows[0].cells] != TABLE_HEADERS:
        issues.append("跟进事项表表头不匹配")
    grid = [int(node.get(qn("w:w"))) for node in table._tbl.tblGrid]
    if grid != TABLE_GRID_DXA:
        issues.append(f"表格列宽不匹配：{grid}")

    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    margins = properties.find(qn("w:tblCellMar"))
    expected_borders = {
        "top": "single", "left": "none", "bottom": "single",
        "right": "none", "insideH": "single", "insideV": "single",
    }
    actual_borders = {
        name: _child_value(borders, f"w:{name}", "w:val") for name in expected_borders
    }
    if actual_borders != expected_borders:
        issues.append(f"表格边框与模板不匹配：{actual_borders}")
    expected_margins = {"top": "0", "left": "108", "bottom": "0", "right": "108"}
    actual_margins = {
        name: _child_value(margins, f"w:{name}", "w:w") for name in expected_margins
    }
    if actual_margins != expected_margins:
        issues.append(f"表格单元格边距与模板不匹配：{actual_margins}")
    if _child_value(properties, "w:tblLayout", "w:type") != "fixed":
        issues.append("表格布局不是模板规定的固定布局")
    if _child_value(properties, "w:tblInd", "w:w") != "0":
        issues.append("表格缩进与模板不匹配")

    for row_index, row in enumerate(table.rows, 1):
        for column_index, cell in enumerate(row.cells, 1):
            if cell.paragraphs[0].alignment != WD_ALIGN_PARAGRAPH.CENTER:
                issues.append(f"表格第{row_index}行第{column_index}列未按模板居中")
            shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
            fill = shading.get(qn("w:fill")) if shading is not None else None
            if fill not in (None, "auto", "FFFFFF"):
                issues.append(f"表格第{row_index}行第{column_index}列存在填充色：{fill}")
    return issues


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx")
    parser.add_argument("--topic", action="append", default=[])
    parser.add_argument("--recorder")
    parser.add_argument("--participant", action="append", default=[])
    args = parser.parse_args()
    issues = validate_minutes(args.docx, args.topic, args.recorder, args.participant)
    if issues:
        for issue in issues:
            print(f"FAIL: {issue}")
        return 1
    print("OK: 标题颜色、第五标题、日期、人员和表格模板格式均符合要求")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
