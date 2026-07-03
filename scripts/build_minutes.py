#!/usr/bin/env python3
"""Build a meeting-minutes DOCX from structured content and the bundled template."""

from __future__ import annotations

import argparse
from copy import deepcopy
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


TABLE_HEADERS = ["序号", "跟进事宜", "负责人", "截止日期", "备注"]
TABLE_GRID_DXA = [817, 4527, 1023, 1295, 1580]
SECTION_NUMERALS = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十"]
TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "assets" / "materials-department-minutes-template.docx"


def today_shanghai() -> date:
    return datetime.now(ZoneInfo("Asia/Shanghai")).date()


def build_title(topics: list[str]) -> str:
    cleaned = [str(item).strip() for item in topics if str(item).strip()]
    if len(cleaned) == 1:
        return f"关于{cleaned[0]}纪要"
    if len(cleaned) <= 3:
        return f"关于{'、'.join(cleaned)}纪要"
    return f"关于{'、'.join(cleaned[:3])}等议题纪要"


def _validate_content(content: dict[str, Any]) -> None:
    topics = content.get("topics") or []
    if not isinstance(topics, list) or not any(str(item).strip() for item in topics):
        raise ValueError("会议主题不能为空")
    if not str(content.get("recorder") or "").strip():
        raise ValueError("纪要人员不能为空")
    participants = content.get("participants") or []
    if not isinstance(participants, list) or not any(str(item).strip() for item in participants):
        raise ValueError("参会人员不能为空")
    for field in ("sections", "conclusions", "actions"):
        if not isinstance(content.get(field, []), list):
            raise ValueError(f"{field} 必须是列表")


def _clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_run_font(run, size: float = 12, bold: bool | None = None, east_asia: str = "宋体") -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), east_asia)


def _format_paragraph(paragraph, *, after: float = 0, line: float = 1.5) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def _add_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Heading 2")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    _set_run_font(paragraph.add_run(text), bold=True)


def _add_subheading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(2)
    _set_run_font(paragraph.add_run(text), bold=True)


def _add_body(doc: Document, text: str, *, indent: bool = True, bold: bool = False) -> None:
    paragraph = doc.add_paragraph()
    _format_paragraph(paragraph, after=3)
    if indent:
        paragraph.paragraph_format.first_line_indent = Pt(24)
    _set_run_font(paragraph.add_run(str(text).strip()), bold=bold)


def _add_labeled_line(doc: Document, label: str, value: str, *, bold_label: bool = True) -> None:
    paragraph = doc.add_paragraph()
    _format_paragraph(paragraph)
    _set_run_font(paragraph.add_run(label), bold=bold_label)
    _set_run_font(paragraph.add_run(value))


def _remove_cell_fill(cell) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    for shading in list(tcpr.findall(qn("w:shd"))):
        tcpr.remove(shading)


def _set_table_geometry(table) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(TABLE_GRID_DXA)))
    tblw.set(qn("w:type"), "dxa")
    tblind = OxmlElement("w:tblInd")
    tblind.set(qn("w:w"), "120")
    tblind.set(qn("w:type"), "dxa")
    tblpr.append(tblind)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in TABLE_GRID_DXA:
        node = OxmlElement("w:gridCol")
        node.set(qn("w:w"), str(width))
        grid.append(node)
    for row in table.rows:
        for cell, width in zip(row.cells, TABLE_GRID_DXA):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")


def _add_actions_table(doc: Document, actions: list[dict[str, Any]]) -> None:
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for index, header in enumerate(TABLE_HEADERS):
        table.rows[0].cells[index].text = header
    for row_index in range(max(4, len(actions))):
        action = actions[row_index] if row_index < len(actions) else {}
        values = [
            str(row_index + 1) if row_index < 2 or row_index < len(actions) else "",
            str(action.get("item", "")),
            str(action.get("owner", "")),
            str(action.get("deadline", "")),
            str(action.get("notes", "")),
        ]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value

    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(marker)
    _set_table_geometry(table)
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            _remove_cell_fill(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
                _format_paragraph(paragraph, line=1.15)
                for run in paragraph.runs:
                    _set_run_font(run, bold=(row_index == 0))


def _find_paragraph(doc: Document, text: str):
    try:
        return next(paragraph for paragraph in doc.paragraphs if paragraph.text == text)
    except StopIteration as exc:
        raise ValueError(f"模板缺少定位标记：{text}") from exc


def _replace_paragraph_text(paragraph, text: str, *, bold: bool | None = None) -> None:
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    for extra in list(paragraph.runs[1:]):
        paragraph._p.remove(extra._r)
    run.text = str(text)
    if bold is not None:
        run.bold = bold


def _remove_paragraph(paragraph) -> None:
    paragraph._p.getparent().remove(paragraph._p)


def _insert_before(marker, text: str, *, bold: bool = False, indent: bool = False):
    paragraph_xml = deepcopy(marker._p)
    marker._p.addprevious(paragraph_xml)
    paragraph = Paragraph(paragraph_xml, marker._parent)
    _replace_paragraph_text(paragraph, text, bold=bold)
    paragraph.paragraph_format.first_line_indent = Pt(24) if indent else None
    return paragraph


def _populate_main_content(doc: Document, sections: list[dict[str, Any]]) -> None:
    marker = _find_paragraph(doc, "{{MAIN_CONTENT}}")
    source = sections or [{"title": "会议议题", "paragraphs": ["相关内容待核验。"]}]
    for index, section in enumerate(source, start=1):
        numeral = SECTION_NUMERALS[index - 1] if index <= len(SECTION_NUMERALS) else str(index)
        title = str(section.get("title") or f"议题{index}").strip()
        _insert_before(marker, f"{numeral}）{title}", bold=True)
        for text in section.get("paragraphs") or ["相关内容待核验。"]:
            _insert_before(marker, str(text).strip(), indent=True)
    _remove_paragraph(marker)


def _populate_conclusions(doc: Document, conclusions: list[Any]) -> None:
    marker = _find_paragraph(doc, "{{CONCLUSIONS}}")
    values = conclusions or ["本次会议未形成可确认的核心结论。"]
    for value in values:
        _insert_before(marker, str(value).strip())
    _remove_paragraph(marker)


def _replace_cell_text(cell, text: str, *, bold: bool | None = None) -> None:
    paragraph = cell.paragraphs[0]
    for extra in list(cell.paragraphs[1:]):
        cell._tc.remove(extra._p)
    _replace_paragraph_text(paragraph, text, bold=bold)
    _remove_cell_fill(cell)


def _fill_actions_table(table, actions: list[dict[str, Any]]) -> None:
    required_data_rows = max(4, len(actions))
    while len(table.rows) - 1 < required_data_rows:
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    for index, header in enumerate(TABLE_HEADERS):
        _replace_cell_text(table.rows[0].cells[index], header, bold=True)
    for index, row in enumerate(table.rows[1:], start=1):
        action = actions[index - 1] if index <= len(actions) else {}
        values = [
            str(index) if index <= 2 or action else "",
            str(action.get("item", "")),
            str(action.get("owner", "")),
            str(action.get("deadline", "")),
            str(action.get("notes", "")),
        ]
        for cell, value in zip(row.cells, values):
            _replace_cell_text(cell, value, bold=False)


def _normalize_review_items(content: dict[str, Any]) -> list[str]:
    items = [str(item).strip() for item in content.get("review_items", []) if str(item).strip()]
    for item in ("专业名称及产品名称", "关键数值", "负责人和截止日期"):
        if len(items) >= 3:
            break
        if item not in items:
            items.append(item)
    return items[:5]


def _format_next_topics(value: Any) -> str:
    if isinstance(value, list):
        return "、".join(str(item).strip() for item in value if str(item).strip()) or "待确认"
    return str(value or "").strip() or "待确认"


def build_minutes(content: dict[str, Any], output_path: str | Path, *, meeting_date: date | None = None) -> Path:
    _validate_content(content)
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"缺少模板资产：{TEMPLATE_PATH}")

    meeting_date = meeting_date or today_shanghai()
    output = Path(output_path).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    topics = [str(item).strip() for item in content["topics"] if str(item).strip()]
    recorder = str(content["recorder"]).strip()
    participants = [str(item).strip() for item in content["participants"] if str(item).strip()]

    doc = Document(TEMPLATE_PATH)
    title_text = build_title(topics)
    source_note = str(content.get("source_note") or "用户提供的会议逐字稿 TXT").strip()
    replacements = {
        "{{TITLE}}": title_text,
        "会议时间：{{MEETING_DATE}}": f"会议时间：{meeting_date.year}年{meeting_date.month}月{meeting_date.day}日",
        "会议地点：{{LOCATION}}": f"会议地点：{str(content.get('location') or '待核验')}",
        "主要事宜：{{TOPICS}}": f"主要事宜：{'、'.join(topics)}",
        "参与人员：{{PARTICIPANTS}}": f"参与人员：{'、'.join(participants)}",
        "纪要人员：{{RECORDER}}": f"纪要人员：{recorder}",
        "下次交流预计时间：{{NEXT_MEETING_DATE}}": f"下次交流预计时间：{str(content.get('next_meeting_date') or '待确认')}",
        "预计议题：{{NEXT_TOPICS}}": f"预计议题：{_format_next_topics(content.get('next_topics'))}",
        "{{IMAGE_NOTE}}": "本次提供的逐字稿未附交流图片。",
        "{{SOURCE_NOTE}}": f"资料来源：{source_note}。未进行外部联网检索。",
        "{{REVIEW_ITEMS}}": "最需要人工复核的信息：" + "；".join(_normalize_review_items(content)) + "。",
    }
    for marker, value in replacements.items():
        _replace_paragraph_text(_find_paragraph(doc, marker), value)
    _populate_main_content(doc, content.get("sections") or [])
    _populate_conclusions(doc, content.get("conclusions") or [])
    _fill_actions_table(doc.tables[0], content.get("actions") or [])

    doc.core_properties.title = title_text
    doc.core_properties.subject = "、".join(topics)
    doc.core_properties.author = recorder
    doc.core_properties.keywords = "会议纪要, 逐字稿"
    doc.save(output)
    return output


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--content", required=True, help="Codex 整理后的 UTF-8 JSON 文件")
    parser.add_argument("--transcript", required=True, help="原始会议逐字稿 TXT，用于来源校验")
    parser.add_argument("--topics", action="append", default=[], help="会议主题，可重复传入")
    parser.add_argument("--recorder", help="纪要人员；提供时覆盖 JSON")
    parser.add_argument("--participant", action="append", default=[], help="参会人员，可重复传入")
    parser.add_argument("--next-meeting-date", help="下次交流预计时间")
    parser.add_argument("--next-topic", action="append", default=[], help="预计议题，可重复传入")
    parser.add_argument("--output", help="输出 DOCX；默认与逐字稿同目录")
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    transcript = Path(args.transcript).resolve()
    if not transcript.is_file():
        raise FileNotFoundError(f"逐字稿不存在：{transcript}")
    try:
        transcript_text = transcript.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("逐字稿必须是 UTF-8 编码") from exc
    if not transcript_text.strip():
        raise ValueError("逐字稿 TXT 为空")

    content = json.loads(Path(args.content).resolve().read_text(encoding="utf-8-sig"))
    if args.topics:
        content["topics"] = [item.strip() for item in args.topics if item.strip()]
    if args.recorder is not None:
        content["recorder"] = args.recorder.strip()
    if args.participant:
        content["participants"] = [item.strip() for item in args.participant if item.strip()]
    if args.next_meeting_date is not None:
        content["next_meeting_date"] = args.next_meeting_date.strip()
    if args.next_topic:
        content["next_topics"] = [item.strip() for item in args.next_topic if item.strip()]
    content.setdefault("source_note", f"会议逐字稿《{transcript.name}》")

    current = today_shanghai()
    output = Path(args.output).resolve() if args.output else transcript.parent / f"{current:%Y%m%d}材料部组会会议纪要.docx"
    print(build_minutes(content, output))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
